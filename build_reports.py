"""Build static research paper and executive summary from verified aggregate outputs."""
from pathlib import Path
from xml.sax.saxutils import escape
import pandas as pd
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, PageBreak, KeepTogether
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_LEFT
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from PIL import Image as PILImage
from matplotlib import font_manager
ROOT=Path(__file__).parent
OUT=ROOT/'outputs'; TABLES=OUT/'tables'; FIGS=OUT/'figures'
pdfmetrics.registerFont(TTFont('DejaVu',font_manager.findfont(font_manager.FontProperties(family='DejaVu Sans'))))
pdfmetrics.registerFont(TTFont('DejaVu-Bold',font_manager.findfont(font_manager.FontProperties(family='DejaVu Sans',weight='bold'))))
pdfmetrics.registerFontFamily('DejaVu',normal='DejaVu',bold='DejaVu-Bold',italic='DejaVu',boldItalic='DejaVu-Bold')
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='TitleX',fontName='DejaVu-Bold',fontSize=23,leading=29,textColor=colors.HexColor('#12354A'),spaceAfter=18))
styles.add(ParagraphStyle(name='HeadX',fontName='DejaVu-Bold',fontSize=15,leading=20,textColor=colors.HexColor('#147D92'),spaceAfter=12))
styles.add(ParagraphStyle(name='BodyX',fontName='DejaVu',fontSize=9.2,leading=14,spaceAfter=8))
styles.add(ParagraphStyle(name='SmallX',fontName='DejaVu',fontSize=7.5,leading=11,spaceAfter=6))
styles.add(ParagraphStyle(name='CellX',fontName='DejaVu',fontSize=7.2,leading=10))
styles.add(ParagraphStyle(name='HeaderCellX',fontName='DejaVu-Bold',fontSize=7.2,leading=10,textColor=colors.white))

def paragraph(x,style='BodyX'):
    return Paragraph(escape(str(x)).replace('\n','<br/>'),styles[style])
def tbl(headers,rows): return ('table',(headers,rows))
def text(x): return ('text',x)
def fig(name,caption): return ('figure',(name,caption))
def formatted(frame,cols=None):
    x=frame if cols is None else frame[cols]
    def f(v):
        if pd.isna(v):return 'N/A'
        if isinstance(v,(int,float)):
            return f'{v:,.0f}' if float(v).is_integer() else f'{v:,.2f}'
        return str(v)
    return [[f(idx)]+[f(v) for v in row] for idx,row in x.iterrows()]

def read(name):return pd.read_csv(TABLES/f'{name}.csv',index_col=0)
age=read('age_summary'); gender=read('gender_summary'); cat=read('category_summary'); levels=read('level_summary')
research=[
('Learner Demographics and Course Enrollment Behavior Analysis on EduPro',[
text('Research paper | Descriptive exploratory data analysis | Prepared 20 October 2026'),
text('Source coverage: 1 January–30 December 2025. Source: the user-supplied EduPro workbook; all findings below are calculated from this file, not from a representative survey.'),
text('Abstract'),
text('This study integrates learner, course and transaction records to describe who enrolls in EduPro courses and how selection varies by demographic group. Validated joins produce 10,000 recorded enrollments from 3,000 learners and 60 courses. Learners aged 26–35 contribute the largest enrollment volume (47.99%), while the under-18 segment has the highest enrollments per registered learner (3.39). Data Science is the largest category (9.16%); category differences are modest. Recorded gender distributions and course-level selection are similar. Activity is concentrated: the top 10% of active learners account for 42.34% of enrollments, while 54% have just one. These are descriptive associations, not evidence of causal demographic effects, completion or skill attainment.'),
tbl(['Core indicator','Verified result'],[
['Registered learners / active learners','3,000 / 3,000'],['Enrollment proxy / catalog courses','10,000 transactions / 60 courses'],['Mean / median enrollments per learner','3.33 / 1'],['Observed age range / mean age','15–35 / 24.97 years'],['Participation within supplied period','100% of registered learners'],['Top 300 learners’ enrollment share','42.34%']]),
text('Study objective'),
text('Provide an auditable baseline for inclusive course planning, learner engagement and outreach. The analysis follows the supplied project brief, emphasizing demographic reach, course selection and participation rather than forecasting, revenue or monetization.'),
text('Keywords: exploratory data analysis; learner demographics; course demand; enrollment concentration; education planning.')]),
('1. Data and methodology',[
text('Inputs and integration. The workbook contains Users (3,000 rows), Courses (60), Transactions (10,000), and Teachers (60). Teachers are outside this project’s scope. The analysis loads user ID, age and gender; course ID, name, category, type and level; and transaction ID, user ID, course ID and date. Personal names and email addresses are deliberately excluded.'),
text('Validation. Required fields have no missing cells. User, course and transaction primary identifiers are unique. All transactions match an existing user and course. Ages are integer values from 15 to 35; dates parse successfully. Transactions → Users and Transactions → Courses are validated many-to-one joins; the joined row count remains 10,000. No repeated learner-course pair occurs, and no rows are dropped or imputed.'),
tbl(['Definition','Operational rule'],[
['Enrollment','One recorded transaction; no attendance/completion claim'],['Active learner','At least one transaction in the supplied period'],['Participation rate','Active learners ÷ registered learners × 100'],['Mean courses per learner','Transactions ÷ registered or active learners (labeled)'],['Category popularity index','Category enrollments ÷ all enrollments × 100'],['Within-group preference','Group-category or group-level count ÷ group enrollments × 100'],['Top 10% concentration','Enrollment share of ceil(0.10 × active learners) highest-count users']]),
text('Age bands. The nonoverlapping groups are <18 (through 17), 18–25, 26–35, 36–45 and 46+. The brief’s “45+” label is replaced with 46+ because age 45 is already in 36–45. Empty groups retain zero counts and undefined (N/A) rates.'),
text('Analytical approach. Compute demographic counts, participation rates, enrollment distributions, within-segment cross-tabulations and concentration measures. Interpret course demand alongside catalog supply. No inferential tests or causal models are used; repeated enrollments within learners are not treated as independent experimental observations.'),
text('Scope of activity. “Active” refers to any qualifying transaction during 2025 in this workbook, not current activity in October 2026. The dashboard recomputes results from a supplied/uploaded workbook; it is interactive analytics, not an external real-time data feed.')]),
('2. Learner demographics and age participation',[
tbl(['Age group','Learners','Enrollments','Enrollments / learner'],formatted(age,['RegisteredLearners','Enrollments','EnrollmentsPerRegisteredLearner'])),
text('The 26–35 segment has 1,446 learners (48.20%) and 4,799 enrollments (47.99%). Ages 18–25 account for 1,121 learners and 3,732 enrollments. The under-18 group contains 433 learners and 1,469 enrollments. No learners occur in either older band; this is a coverage gap, not evidence that older adults lack interest.'),
text('Enrollment intensity differs little: 3.39 transactions per learner for <18, 3.33 for 18–25 and 3.32 for 26–35. All groups with learners have 100% participation in the supplied period. Thus, highest volume and highest per-learner intensity identify different groups.'),
fig('02_age_participation','Figure 1. Registered learners, enrollment volume and per-learner enrollment rate by age band.'),
text('Recorded gender includes 1,520 Female learners (50.67%) and 1,480 Male learners (49.33%). Female learners contribute 5,078 transactions (50.78%) and Male learners 4,922 (49.22%). Enrollments per learner are 3.34 and 3.33, respectively. The workbook contains only these two recorded labels, which limits conclusions about broader gender inclusivity.')]),
('3. Course category popularity',[
tbl(['Category','Enrollments','Share (%)','Enrollments / course'],formatted(cat,['Enrollments','EnrollmentSharePct','EnrollmentsPerCatalogCourse'])),
text('Data Science leads with 916 enrollments (9.16%), followed by Finance (864; 8.64%) and Web Development (844; 8.44%). Marketing and Programming tie for the lowest count (806; 8.06% each). All 12 categories contain five courses, making per-category catalog supply equal. The largest-to-smallest share difference is only 1.10 percentage points; rank alone should not be interpreted as a large demand divide.'),
fig('03_category_popularity','Figure 2. Recorded category enrollment totals; all categories contain five courses.'),
text('Implication. Review capacity and learner feedback in Data Science first, but do not discontinue lower-ranked categories on this evidence alone. Even the lowest-ranked categories account for more than 800 recorded enrollments each.')]),
('4. Demographics × course preference',[
text('Age and category. Finance is the leading category among under-18 enrollments (9.26%; 136 transactions). Data Science leads ages 18–25 (9.19%; 343) and 26–35 (9.56%; 459). Each represented age segment distributes its enrollments across all categories. Row-normalized shares are used to avoid mistaking larger population size for stronger category preference.'),
fig('05_age_category_heatmap','Figure 3. Within-age category enrollment shares (%). Empty older groups have no estimable preference.'),
text('Gender and category. Data Science is the top category for both Female (9.43%) and Male (8.88%) enrollments. The largest absolute difference between gender-specific category shares is Marketing: Male 8.55% versus Female 7.58%, a gap of 0.97 percentage points. All observed category-share gaps are below one percentage point; no statistical significance or causal gender effect is claimed.'),
fig('06_demographics_level','Figure 4. Within-group course-level enrollment shares by recorded gender and age.'),
text('Gender and level. Beginner-course shares are 35.66% for Female and 35.80% for Male enrollments; Advanced shares are 34.62% and 34.88%. Under-18 learners select Advanced courses most often by a small margin (36.35% versus 34.10% Beginner). Beginner is the leading level for ages 18–25 (36.04%) and 26–35 (35.99%). These are metadata-based selections, not measured skill differences.')]),
('5. Course type, level and beginner-course behavior',[
tbl(['Course level','Catalog courses','Enrollments','Share (%)'],formatted(levels,['CatalogCourses','Enrollments','EnrollmentSharePct'])),
text('Beginner courses attract 3,573 enrollments (35.73%), Advanced 3,475 (34.75%), and Intermediate 2,952 (29.52%). Catalog supply is 21 Beginner, 21 Advanced and 18 Intermediate courses. Enrollment per catalog course is 170.14, 165.48 and 164.00, respectively. The lower Intermediate total partly reflects fewer available courses, rather than establishing a substantial preference gap.'),
text('Free courses account for 6,403 enrollments (64.03%), compared with 3,597 Paid enrollments (35.97%). The catalog has 38 Free courses and 22 Paid courses: average enrollments per course are 168.50 and 163.50. This is not proof that price causes demand differences because visibility, content, access and other factors are not controlled.'),
tbl(['Course level','Free enrollments','Paid enrollments','Free share (%)'],[
['Beginner','2,395','1,178','67.03'],['Intermediate','2,003','949','67.85'],['Advanced','2,005','1,470','57.70']]),
text('Within Beginner offerings, 14 of 21 courses are Free (66.67% of supply), closely matching the 67.03% enrollment share. Free and Paid Beginner courses average 171.07 and 168.29 enrollments per course. The observed split therefore does not establish a strong learner preference beyond availability.'),
text('Overlapping learner cohorts. A total of 1,596 learners select at least one Beginner course and 1,608 select at least one Advanced course; 752 select both. Mean same-level enrollments per selecting learner are 2.24 and 2.16, respectively. The groups overlap and must not be treated as separate, verified beginner versus advanced skill populations.'),
text('Metadata quality. “Advanced Python” is labeled Beginner in the source. The analysis preserves the recorded CourseLevel rather than silently relabeling it. A subject-matter review is needed before using levels for learning pathways or communications.')]),
('6. Enrollment concentration and activity distribution',[
text('The mean is 3.33 enrollments per learner but the median is one. A total of 1,620 learners (54.00%) have one enrollment and 612 (20.40%) have two. This difference between mean and median shows that a high-activity minority raises the overall average.'),
tbl(['Top active learners (%)','Learners','Enrollments','Enrollment share (%)'],formatted(read('enrollment_concentration'),['Learners','Enrollments','EnrollmentSharePct'])),
fig('09_activity_concentration','Figure 5. Activity-count distribution and cumulative share of enrollments (highest-activity learners first).'),
text('The top 300 active learners generate 4,234 enrollments (42.34%). This supports monitoring both low-frequency participants and highly active learners rather than treating the mean as a typical learner journey.'),
text('A notable structural pattern is the absence of learners with five to eight enrollments; observed counts are 1–4 and 9–16. This is a verified distributional feature, but its cause is unknown. It could reflect data construction, program rules or other unobserved factors. Validate data provenance before generalizing to real platform behavior.'),
text('A single enrollment is not evidence of churn or dissatisfaction, and many enrollments are not evidence of course completion or successful learning. Transaction timing, course requirements and outcomes would be needed to evaluate those interpretations.')]),
('7. Recommendations, limitations and conclusion',[
text('Recommendations for platform and government stakeholders'),
text('1. Establish an inclusive reporting baseline. Publish counts, within-group enrollment shares and per-learner rates together. Report missing age coverage and the limited recorded gender labels. Do not describe participation parity as evidence of equal access or equal learning outcomes.'),
text('2. Pilot demand-informed capacity planning. Review Data Science capacity and learner feedback first while maintaining the broader catalog. Treat modest category differences as a prioritization signal, not grounds for major resource withdrawal. Compare demand per offered course.'),
text('3. Investigate low-frequency participation. Invite feedback from one-enrollment learners and assess whether one course met their goals. Test onboarding, navigation and course-discovery improvements; evaluate using subsequent participation and completion rather than assuming disengagement.'),
text('4. Validate taxonomy and data provenance. Review title/level inconsistencies and the activity-count gap. Document how the workbook was collected or generated. Maintain validated primary keys, foreign keys and age/date rules at each refresh.'),
text('5. Improve evidence for equitable education planning. Collect optional learning goals, self-assessed prior experience, accessibility requirements, appropriate inclusive demographics and completion/outcome indicators with consent and data minimization. Investigate older learner representation before designing or targeting an older-adult program.'),
text('6. Keep decisions proportionate to evidence. Use small, measurable pilots instead of demographic stereotyping. Course-level labels must not be used as learner-skill labels. Apply privacy safeguards to minor learners and use aggregate reporting with an appropriate small-cell policy before wider publication.'),
text('Limitations'),
text('This is one supplied workbook covering transactions in 2025, with ages restricted to 15–35. Data provenance and representativeness are unverified. Enrollment is a transaction proxy; no attendance, completion, learning goal, skill or outcome measure is available. Course type and level vary with catalog supply, and metadata inconsistencies remain. There is no external benchmark or causal identification. These constraints limit generalization to platform growth, public education access and learner success.'),
text('Conclusion'),
text('EduPro’s recorded demand is broadly distributed across categories and genders, with 26–35 learners contributing the most enrollment volume and a high-activity minority accounting for a large share of transactions. The strongest immediate opportunities are better taxonomy and provenance checks, improved participation measurement, inclusive data collection and carefully evaluated learner-support pilots.')]),
('8. Reproducibility and deliverable guide',[
text('The executed notebook contains ten numbered stages, with Python code and actual printed tables or charts directly beneath each stage. The companion Python script contains the same analysis; the shared data utility module exposes all loading, validation, integration and denominator logic.'),
tbl(['Step','Code activity','Observed output'],[
['1','Import libraries; create output folders','Ready message'],['2','Load workbook, audit IDs/fields; validated joins','3,000 users; 60 courses; 10,000 joined rows; audit table'],['3','Define KPIs and concentration denominator','KPI table: mean 3.33, median 1, top-10% share 42.34%'],['4','Group demographics and age participation','Age/gender summary tables; two chart panels'],['5','Group category/type/level demand and catalog supply','Demand tables; category and type/level charts'],['6','Cross-tab age × category; row-normalize','Count and share tables; heatmap'],['7','Cross-tab gender preferences and age × level','Six tables, share-gap output and two chart panels'],['8','Compare course levels, types and learner overlap','Supply-adjusted tables; 752 overlapping learners'],['9','Compute activity frequency and concentration','Frequency/concentration tables and chart panel'],['10','Export workbook and print conclusions','Aggregate Excel workbook and verified findings']]),
text('Execution. Install the supplied requirements, place the original workbook beside the source files, then run the analysis script or open and run the notebook. Start the Streamlit app separately. The app also accepts an uploaded copy of the workbook. Detailed commands and deployment notes are supplied in the README.'),
text('Dashboard coverage. Learner overview, age-wise enrollment, gender-based course preferences, category/type/level popularity, activity concentration and data quality. Filters: age group, gender, category, course level and course type. All outputs recalculate for the selection; empty filters and empty age groups are handled explicitly. Registered learner denominators follow demographic filters only.'),
text('Verification. The executed notebook has no error outputs. Streamlit automated smoke tests cover unfiltered data, an empty older-age group, a selected age/category combination, and an empty category selection. These checks validate code execution and denominator handling, not every possible user interaction.'),
text('Source material and references. (1) User-supplied EduPro Online Platform workbook, sheets Users, Courses and Transactions. (2) User-supplied project brief. No external substantive evidence was introduced. All tables and figures are original computations from the supplied data.'),
text('Privacy. Only aggregate tables and charts are exported. Raw learner identifiers, names and email addresses are not included in the submission package. The original workbook must be supplied locally or uploaded to run a fresh analysis. Government stakeholder use is a suggested audience, not evidence that the workbook represents a government program.')])]

executive=[('EduPro | Executive summary',[
text('For education and government stakeholders | Evidence window: 1 January–30 December 2025'),
text('Purpose. Describe who enrolls and what courses they select, using the supplied workbook. This is not an evaluation of learning outcomes, program impact or population-level access.'),
tbl(['Headline indicator','Finding'],[
['Participation baseline','3,000 registered learners; all have at least one transaction'],['Recorded enrollment','10,000 transactions across 60 courses'],['Largest age segment','26–35: 47.99% of enrollment; 48.20% of learners'],['Highest enrollment rate','Under 18: 3.39 enrollments per learner'],['Gender distribution','Female: 50.78% of enrollment; Male: 49.22%'],['Top category','Data Science: 916 enrollments (9.16%)'],['Enrollment intensity','Mean 3.33; median 1; 54% have one enrollment'],['Concentration','Top 300 learners account for 42.34% of enrollment']]),
text('What the evidence suggests'),
text('Demand is broad rather than dominated by one subject: category shares range from 8.06% to 9.16%. Gender-based category-share gaps are all below one percentage point. Most of the age-volume difference reflects the number of learners in each band, while per-learner enrollment rates are similar.'),
text('Beginner courses lead narrowly (35.73%), followed by Advanced (34.75%). Free courses account for 64.03% of enrollment, but also 63.33% of the catalog. Course availability should be considered before attributing these totals to learner preference. “Beginner” is a course label, not a verified learner skill.'),
text('Recommended decisions'),
text('• Prioritize a review of Data Science capacity and feedback without removing other subjects on modest rank differences.\n• Investigate whether one-enrollment learners achieved their goals before treating them as disengaged.\n• Review course-level metadata and the unusual activity-count distribution.\n• Add optional learning goals, prior experience, accessibility and completion measures before evaluating learner success.\n• Report reach, intensity and outcomes separately; use measured pilots rather than demographic assumptions.'),
text('Evidence cautions'),
text('Ages are restricted to 15–35, and only Female/Male gender labels appear. The source’s representativeness is unknown. A transaction does not prove attendance, completion or skill attainment. No causal or statistical-significance claim is made. One source labels Advanced Python as Beginner. No learners have five to eight enrollments; provenance should be checked.'),
text('Delivery. The accompanying executed notebook, aggregate workbook, research paper and interactive Streamlit source make the analysis reproducible. The dashboard must be launched locally or deployed; it is not a hosted live service.')])]

def footer(canvas,doc):
    canvas.saveState(); canvas.setFont('DejaVu',7)
    canvas.setFillColor(colors.HexColor('#687A86'))
    canvas.drawString(42,26,'EDUPRO  |  Descriptive learner intelligence  |  Source: supplied workbook')
    canvas.drawRightString(553,26,str(doc.page)); canvas.restoreState()

def build(pages,stem):
    story=[]; document=Document()
    for sec in document.sections:
        sec.top_margin=Inches(.65);sec.bottom_margin=Inches(.65)
    normal=document.styles['Normal'];normal.font.name='Calibri';normal.font.size=Pt(10)
    for page_i,(heading,items) in enumerate(pages):
        if page_i: story.append(PageBreak());document.add_page_break()
        story.append(paragraph(heading,'TitleX' if page_i==0 else 'HeadX'))
        document.add_heading(heading,0 if page_i==0 else 1)
        for kind,item in items:
            if kind=='text':
                story.append(paragraph(item))
                document.add_paragraph(item)
            elif kind=='table':
                headers,rows=item
                data=[[paragraph(h,'HeaderCellX') for h in headers]]+[[paragraph(x,'CellX') for x in r] for r in rows]
                n=len(headers)
                widths=[511/n]*n
                if n==2:widths=[178,333]
                if n==4 and headers[0]=='Category':widths=[190,100,100,121]
                if n==3 and headers[0]=='Step':widths=[36,220,255]
                table=Table(data,colWidths=widths,repeatRows=1,hAlign='LEFT')
                table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#147D92')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.HexColor('#EFF5F7'),colors.white]),('VALIGN',(0,0),(-1,-1),'TOP'),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('LINEBELOW',(0,-1),(-1,-1),.5,colors.HexColor('#C9D8DF'))]))
                story += [table,Spacer(1,10)]
                t=document.add_table(rows=1,cols=n);t.style='Light Shading Accent 1'
                for cell,h in zip(t.rows[0].cells,headers):cell.text=str(h)
                for row in rows:
                    for cell,v in zip(t.add_row().cells,row):cell.text=str(v)
                document.add_paragraph('')
            else:
                name,caption=item;path=FIGS/f'{name}.png'
                with PILImage.open(path) as im:w,h=im.size
                width=505;height=width*h/w
                story.append(KeepTogether([Image(str(path),width=width,height=height),paragraph(caption,'SmallX')]))
                document.add_picture(str(path),width=Inches(6.35))
                document.add_paragraph(caption,style='Caption')
    pdf=OUT/f'{stem}.pdf';docx=OUT/f'{stem}.docx'
    SimpleDocTemplate(str(pdf),pagesize=(595,842),rightMargin=42,leftMargin=42,topMargin=40,bottomMargin=43).build(story,onFirstPage=footer,onLaterPages=footer)
    document.save(docx)
    print(pdf);print(docx)
build(research,'EduPro_research_paper')
build(executive,'EduPro_executive_summary')
