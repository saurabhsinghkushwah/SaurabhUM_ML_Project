from pathlib import Path
import numpy as np
import pandas as pd
from analysis_core import kpis, summarize, customer_table, discount_scenario

def fmt(v): return f'{v:,.2f}'
def table_md(d):
    x = d.copy()
    for c in x.select_dtypes(include='number'):
        x[c] = x[c].map(lambda v: f'{v:,.2f}')
    return '| ' + ' | '.join(x.columns) + ' |\n|' + '|'.join(['---'] * len(x.columns)) + '|\n' + '\n'.join('| ' + ' | '.join(map(str, row)) + ' |' for row in x.itertuples(index=False, name=None))

def word_copy(text, path):
    from docx import Document
    doc = Document()
    for line in text.splitlines():
        if line.startswith('# '): doc.add_heading(line[2:], 0)
        elif line.startswith('## '): doc.add_heading(line[3:], 1)
        elif line.startswith('### '): doc.add_heading(line[4:], 2)
        elif line.startswith('- '): doc.add_paragraph(line[2:], style='List Bullet')
        elif line.strip(): doc.add_paragraph(line)
    doc.save(path)

def write_reports(d, audit, rejected, out):
    out = Path(out)
    out.mkdir(exist_ok=True)
    k = kpis(d)
    cats = summarize(d, ['Category Name'])
    products = summarize(d, ['Product Name'])
    markets = summarize(d, ['Market'])
    segments = summarize(d, ['Customer Segment'])
    statuses = summarize(d, ['Order Status'])
    bands = summarize(d, ['discount_band']).sort_values('discount_band')
    cs = customer_table(d)
    losses = cs[cs['profit'] < 0]
    pos = cs[cs['profit'] > 0]
    n = max(1, int(np.ceil(.2 * len(pos))))
    concentration = 100 * pos.head(n)['profit'].sum() / pos['profit'].sum()
    loss_cats = cats[cats['profit'] < 0]
    loss_products = products[products['profit'] < 0]
    weak = markets[(markets['net_revenue'] >= markets['net_revenue'].median()) & (markets['margin_pct'] < k['margin_pct'])]
    finalized = kpis(d[d['Order Status'].isin(['COMPLETE', 'CLOSED'])])
    cols = ['net_revenue', 'profit', 'margin_pct']
    s0 = discount_scenario(d, 0)
    corr = d['discount_rate'].corr(d['margin_pct'], method='spearman')
    report = f'''# Customer, Product, and Profitability Performance Analysis in Supply Chain Operations
## Research paper: analysis of the supplied APL Logistics case dataset

## Abstract
This descriptive study evaluates {len(d):,} retained records and {k['customers']:,} customer identifiers from the supplied CSV. Gross sales total {fmt(k['gross_sales'])}, net revenue {fmt(k['net_revenue'])}, and reported profit {fmt(k['profit'])}, producing a weighted net margin of {k['margin_pct']:.2f}%. All monetary amounts are dataset currency units; no currency code was provided. The analysis identifies {len(losses):,} loss-making customers, {len(loss_products):,} loss-making products and {len(loss_cats):,} loss-making categories. Results support prioritization and further validation, not audited revenue recognition or causal pricing conclusions.

## 1. Research questions and scope
Which customers, products, categories and destination markets contribute profit rather than merely sales? How are observed discounts associated with reported margins? What does a fixed-cost pricing scenario imply? This study follows the supplied project requirements. The entity label comes from the assignment, and the file's provenance and accounting completeness have not been independently verified. The population is the supplied file, not necessarily the provider's actual operations.

## 2. Data and financial validation
The input contains 40 columns. It has no order date, shipping date, unique order or line identifier, currency identifier, exchange rate, or explicit shipping-cost field. The analysis therefore treats observations as records/apparent item-level economics, does not report distinct orders, and does not construct time trends. Shipping-cost erosion cannot be isolated.

The CSV is read using UTF-8 with a Windows-1252 fallback. Text is trimmed; numeric conversion is explicit. Invalid/missing financial values, nonpositive revenue, invalid quantities, invalid discount rates and inconsistent financial aliases or customer identifiers are quarantined. A two-cent absolute reconciliation tolerance is used for rounded financial fields. Retained records total {len(d):,}; quarantined records total {len(rejected):,}. Exact duplicate rows, if encountered, are not automatically removed because identity cannot be established without record keys. Negative profit is retained because loss diagnosis is central to the assignment. Nonfinancial missing names or postal codes do not justify removing financially valid rows.

Sales reconciles to quantity times item price. Net revenue reconciles to Sales minus the discount amount. Benefit per order duplicates Order Profit Per Order, while Sales per customer duplicates Order Item Total. The aliases are not added together. The recorded discount rate is rounded and can differ from discount amount divided by gross sales; monetary reconciliation uses the recorded discount amount. All order statuses remain in the primary scope, with a separate COMPLETE/CLOSED sensitivity proxy.

## 3. Methods and KPI definitions
Gross revenue is the sum of Sales, satisfying the brief's total-revenue definition. Net revenue is the sum of Order Item Total. Profit is the sum of Order Profit Per Order once. The primary margin is total profit divided by total net revenue, multiplied by 100. Profit divided by gross sales is reported separately ({k['profit_to_gross_pct']:.2f}%). Category and market margins are revenue-weighted, not unweighted averages of row ratios.

Customer Value Index is explicitly defined here as each customer's aggregate reported profit; it is a project-specific definition rather than a standardized external index. Positive-profit customers at or above their 80th-percentile profit cutoff form the high-value tier. Other profitable customers form the positive core; customers with zero or negative aggregate profit form the nonpositive tier. A secondary low-margin flag uses positive profit with margin below 5%. These are transparent managerial thresholds rather than learned optimal boundaries.

Product prioritization uses the top quartile of product revenue combined with the bottom quartile of product margin. Strong-revenue/weak-margin markets have revenue at or above the market median and margin below the full portfolio. Discount bands and rank correlations are descriptive. Product/market stratification with at least 30 records per cell reduces some mix effects but is not causal identification. No significance tests or causal confidence intervals are asserted.

## 4. Revenue and reported profit results
Gross sales: {fmt(k['gross_sales'])}; discounts: {fmt(k['discounts'])}; net revenue: {fmt(k['net_revenue'])}; reported profit: {fmt(k['profit'])}. The primary weighted net margin is {k['margin_pct']:.2f}%.

Restricting to COMPLETE/CLOSED records yields net revenue {fmt(finalized['net_revenue'])}, reported profit {fmt(finalized['profit'])}, and margin {finalized['margin_pct']:.2f}%. This demonstrates sensitivity to order-status scope, not verified recognized revenue. Pending, canceled and suspected-fraud records should be reconciled with the accounting owner before external reporting. Revenue and margin time trends cannot be assessed in the supplied file.

### Status breakdown
{table_md(statuses[['Order Status'] + cols])}

## 5. Customer contribution
There are {len(losses):,} customers with negative aggregate profit; their combined reported loss is {fmt(losses['profit'].sum())}. The top 20% of positive-profit customers account for {concentration:.2f}% of the positive-profit pool. This concentration metric uses positive profits as the denominator so losses do not create misleading contribution percentages. It is not a percentage of net portfolio profit. Customer IDs in detailed outputs require restricted access and should not appear in public executive reports.

### Segment economics
{table_md(segments[['Customer Segment'] + cols])}

## 6. Product and category profitability
The highest-profit category is {cats.iloc[0]['Category Name']} with reported profit {fmt(cats.iloc[0]['profit'])} and margin {cats.iloc[0]['margin_pct']:.2f}%. The lowest-profit category is {cats.iloc[-1]['Category Name']} with profit {fmt(cats.iloc[-1]['profit'])}. There are {len(loss_cats):,} negative-profit categories; a low-margin category should not be mislabeled as loss-making when its total profit is positive.

### Top categories by profit
{table_md(cats[['Category Name'] + cols].head(10))}

### Bottom products by reported profit
{table_md(products[['Product Name'] + cols].nsmallest(5, 'profit'))}

## 7. Discount diagnostics and modeled scenarios
The Spearman association between recorded discount rate and record-level net margin is {corr:.4f}. This is observational: discounts may vary with product, customer, geography, transaction status and underlying costs. Different observed group margins cannot be interpreted as the profit caused by discounts.

### Observed discount-band economics
{table_md(bands[['discount_band', 'rows'] + cols])}

Discount Impact Ratio is defined as discounts divided by (reported profit plus discounts), multiplied by 100: {k['discount_impact_ratio_pct']:.2f}%. It measures the modeled share of no-discount profit removed by recorded discounts, assuming unchanged costs, quantities and prices. It is not the observed margin difference between discounted and non-discounted groups. The corresponding modeled margin erosion is {k['modeled_margin_erosion_pp']:.2f} percentage points.

A zero-discount scenario gives modeled profit {fmt(s0['scenario_profit'])}. The apparent uplift restores the gross-to-net difference mechanically; it is not a sales forecast. The aggregate residual between gross sales minus recorded discounts and reported net revenue is {fmt(k["gross_sales"] - k["discounts"] - k["net_revenue"])} currency units. All individual rows reconcile within the two-cent tolerance, but accumulated rounding means recorded discounts and gross-to-net differences are not identical. The scenario uses reported net revenue and inferred costs; the Discount Impact Ratio uses recorded discount amounts. For each scenario, modeled net revenue equals gross sales times one minus the replacement discount rate. Modeled profit equals this revenue less implied cost, where implied cost equals observed net revenue minus reported profit. Costs are inferred residuals, not independently measured logistics or shipping costs. Demand response, competitor actions, capacity and customer retention are excluded. No universal causal discount threshold can be justified by these records.

## 8. Market and regional diagnostics
{table_md(markets[['Market'] + cols])}

Markets meeting the stated high-revenue/weak-margin screening rule: {', '.join(weak['Market'].astype(str)) if len(weak) else 'none'}. Geography refers to order destination, not the customer's registered country. Country labels are preserved as supplied because standardization needs a reviewed reference mapping. Region and country summaries are exported for more granular review.

## 9. Recommendations and implementation
- Protect profitable high-value relationships with service-level reviews, while avoiding blanket discounts justified only by revenue.
- Investigate loss-making customers and products for cost allocation, returns, fraud/status effects and unusual pricing before withdrawing service or delisting products.
- Review high-revenue/low-margin products and weak-margin markets using the explicit screening rules; validate unit economics against the ledger.
- Pilot discount guardrails in a controlled experiment, accounting for product mix, demand, costs and retention. Do not use descriptive bands as proven optimal price thresholds.
- Use the dashboard to inspect customer tiers, category/market heatmaps, order-status sensitivity and fixed-cost pricing scenarios.
- Obtain dates, order/line keys, currency codes, exchange rates, shipping costs, returns, taxes, and approved recognition rules before extending to time trends or accounting-grade reporting.

## 10. Government-stakeholder governance and limitations
Public-facing reporting should use aggregated markets/categories rather than customer identifiers, names, addresses or coordinates. Restrict access, review retention periods, and confirm dataset licensing and provenance. Findings are a case-study diagnostic rather than evidence of government spending performance or actual corporate financial outcomes. No taxpayer savings, policy impact, causality or currency denomination is inferred. Generalizability beyond the supplied dataset is unknown.

## Conclusion
Sales alone can conceal weak customer and product economics. This reproducible workflow separates gross sales, net revenue and reported profit, retains losses and explains discount assumptions. Further business action depends on validating costs, recognition policies and source-system metadata.

## Sources and reproducibility
1. User-supplied project brief (paste.txt): scope, fields, KPIs and dashboard requirements.
2. User-supplied APL_Logistics (1).csv: all empirical calculations.
3. Accompanying executed notebook, shared analysis functions, validation exports and dashboard source: computational methods.
No external company claims or literature findings are introduced. This is a reproducible analytical case-study paper, not a systematic literature review.
'''
    executive = f'''# Executive summary for government stakeholders
## Scope and evidence
This case-study analysis uses {len(d):,} retained records and {k['customers']:,} customer identifiers from the supplied logistics dataset. Results are reported dataset economics, not audited corporate statements or verified government-program outcomes. Currency is unspecified.

## Key findings
- Gross sales: {fmt(k['gross_sales'])}; recorded discounts: {fmt(k['discounts'])}.
- Net revenue: {fmt(k['net_revenue'])}; reported profit: {fmt(k['profit'])}; weighted net margin: {k['margin_pct']:.2f}%.
- {len(losses):,} customers have negative aggregate reported profit, totaling {fmt(losses['profit'].sum())}.
- Top 20% of positive-profit customers contribute {concentration:.2f}% of the positive-profit pool.
- Loss-making products: {len(loss_products):,}; loss-making categories: {len(loss_cats):,}.
- Highest-profit category: {cats.iloc[0]['Category Name']}; highest-profit market: {markets.iloc[0]['Market']}.
- COMPLETE/CLOSED sensitivity margin: {finalized['margin_pct']:.2f}%, versus {k['margin_pct']:.2f}% across all statuses.

## Recommended actions
- Validate accounting recognition and costs before using these figures for procurement, budgets or performance claims.
- Review loss-making customers/products and high-revenue/weak-margin markets; avoid indiscriminate revenue-driven discounting.
- Trial pricing changes with controlled evaluation and safeguards for service quality and retention.
- Require source-system dates, currency codes, unique transaction keys and shipping costs for stronger oversight.

## Controls and caveats
All statuses are included unless explicitly filtered; canceled, pending and suspected-fraud records are not asserted to be recognized revenue. No dates exist for trend analysis. Discount associations are not causal. What-if scenarios fix costs and quantities and do not predict demand. No actual government fiscal impact is measured. Publish only aggregate findings, and restrict customer-level outputs. The dashboard is provided as runnable source, not as an already-hosted public service.
'''
    for name, text in [('research_paper', report), ('executive_summary', executive)]:
        (out / f'{name}.md').write_text(text)
        word_copy(text, out / f'{name}.docx')
    print('Verified figures:', {key: round(float(value), 4) for key, value in k.items()})
    print('Loss customers:', len(losses), 'Loss products:', len(loss_products), 'Loss categories:', len(loss_cats))
